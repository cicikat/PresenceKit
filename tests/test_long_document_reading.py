import io
import asyncio
from unittest.mock import AsyncMock

import pytest

from core import character_document_library as library, document_reading, media_processor
from core.tools.character_recall import read_document_for_user
from core.tools.tool_result import to_tool_result, frame_tool_message


def store(text, name="story.md", digest="b" * 64):
    return library.store_upload(uid="reader", char_id="char-test", filename=name,
                                media_type="text/markdown", sha256=digest,
                                searchable_text=text, source="upload_file")


@pytest.mark.asyncio
async def test_ten_thousand_character_upload_can_read_every_character_without_second_clipping(sandbox):
    text = "# 开始\n" + "内容甲乙丙丁" * 2000 + "\n# 结尾\n末尾证据"
    result = await media_processor.ingest_file_bytes(text.encode(), "story.md", uid="reader", char_id="char-test")
    assert result[0] == text
    digest = media_processor._hash_bytes(text.encode())
    doc = "doc_" + digest[:24]
    projection = media_processor.document_upload_context(text, "story.md", digest, uid="reader", char_id="char-test")
    assert doc in projection and str(len(text)) in projection
    assert "本轮先提供前 3000 字符" in projection
    assert "末尾证据" not in projection
    chunks, offset = [], 0
    while True:
        page = library.read("reader", "char-test", doc, offset=offset, limit=6000)
        tool = await read_document_for_user("reader", "char-test", doc, offset=offset)
        safe = to_tool_result(tool).safe_summary
        assert page["content"] in frame_tool_message(safe)
        assert "工具结果已截断" not in safe and len(safe) < 8000
        chunks.append(page["content"])
        offset = page["next_offset"]
        if offset is None:
            break
    assert "".join(chunks) == text
    reopened = library.read("reader", "char-test", doc, mode="continue")
    assert reopened["content"] == ""
    assert reopened["progress"]["provided_complete"]
    assert library.read("reader", "different", doc) is None
    assert library.read("other", "char-test", doc) is None


@pytest.mark.asyncio
async def test_nonsequential_reads_do_not_claim_whole_document_and_resume_first_gap(sandbox):
    doc = store("# 第一章\n" + "a" * 8000 + "\n# 第二章\n" + "b" * 8000)
    await read_document_for_user("reader", "char-test", doc, offset=8000)
    assert library.read("reader", "char-test", doc, mode="continue")["offset"] == 0
    await read_document_for_user("reader", "char-test", doc, limit=1000)
    assert library.read("reader", "char-test", doc, mode="continue")["offset"] == 1000
    outline = library.read("reader", "char-test", doc, mode="overview")
    assert outline["sections"][1]["offset"] > 8000
    assert not outline["progress"]["provided_complete"]


@pytest.mark.asyncio
async def test_summary_checkpoints_retry_cover_tail_and_survive_reload(sandbox, monkeypatch):
    from core import llm_client
    text = "a" * 10000 + "b" * 10000 + "尾部重要事实"
    doc = store(text)
    summarize = AsyncMock(side_effect=["首段概要", None, "首段合并概要"])
    monkeypatch.setattr(llm_client, "summarize_document_text", summarize)
    await document_reading.summarize("reader", "char-test", doc)
    state = library.get_record("reader", "char-test", doc)["overview"]
    assert state["status"] == "partial" and state["covered_chars"] == 10000
    assert len(state["parts"]) == 1
    later = AsyncMock(side_effect=["第二段概要", "尾部重要事实", "含首中尾的全文概要"])
    monkeypatch.setattr(llm_client, "summarize_document_text", later)
    await document_reading.summarize("reader", "char-test", doc)
    assert later.call_args_list[0].args[0] == "b" * 10000
    assert later.call_args_list[1].args[0] == "尾部重要事实"
    state = library.get_record("reader", "char-test", doc)["overview"]
    assert state["status"] == "ready" and state["covered_chars"] == len(text)
    assert state["parts"][-1]["start"] == 20000
    assert library.read("reader", "char-test", doc, mode="summary")["content"] == "含首中尾的全文概要"
    # A duplicate upload must preserve progress and derived overviews.
    assert store(text) == doc
    assert library.get_record("reader", "char-test", doc)["overview"]["status"] == "ready"
    stats = library.observability("reader", "char-test")
    assert stats["documents"][0]["overview_status"] == "ready"
    assert "全文概要" not in str(stats)


@pytest.mark.asyncio
async def test_summary_merge_failure_retries_without_regenerating_parts_or_reviving_deleted_document(sandbox, monkeypatch):
    from core import llm_client
    doc = store("fact" * 3000)
    initial = AsyncMock(side_effect=["first", "second", None])
    monkeypatch.setattr(llm_client, "summarize_document_text", initial)
    await document_reading.summarize("reader", "char-test", doc)
    assert library.get_record("reader", "char-test", doc)["overview"]["covered_chars"] == 0

    async def withdrawn(content, *, combine=False, char_id=None):
        assert combine
        library.delete("reader", "char-test", doc)
        return "merged"

    monkeypatch.setattr(llm_client, "summarize_document_text", withdrawn)
    await document_reading.summarize("reader", "char-test", doc)
    assert library.read("reader", "char-test", doc) is None


def test_docx_headings_and_tables_are_extracted_in_document_order():
    from docx import Document
    doc = Document()
    doc.add_heading("章节一", level=1)
    doc.add_paragraph("表格之前")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "表内重要信息"
    table.cell(0, 1).text = "另一列"
    doc.add_paragraph("表格之后")
    out = io.BytesIO()
    doc.save(out)
    text = media_processor.parse_file_bytes(out.getvalue(), "story.docx")
    assert text.index("表格之前") < text.index("表内重要信息") < text.index("表格之后")
    assert document_reading.sections(text)[0]["title"] == "章节一"
    assert media_processor.parse_file_bytes(b"old binary", "old.doc") is None


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "gbk"])
def test_text_encoding_preserves_complete_content(encoding):
    text = "第一章 开始\n正文\n最后一行"
    assert media_processor.parse_file_bytes(text.encode(encoding), "story.txt") == text


@pytest.mark.asyncio
async def test_text_limit_and_archive_failure_are_explicit(sandbox, monkeypatch):
    with pytest.raises(media_processor.MediaIngestError, match="没有静默截断"):
        await media_processor.ingest_file_bytes(b"x" * 500001, "large.txt", uid="reader", char_id="char-test")
    monkeypatch.setattr(library, "store_upload", lambda **kwargs: None)
    with pytest.raises(media_processor.MediaIngestError) as error:
        await media_processor.ingest_file_bytes(b"content", "story.txt", uid="reader", char_id="char-test")
    assert error.value.code == "document_archive_failed"


@pytest.mark.asyncio
async def test_document_progress_and_overview_continue_without_reannouncing_upload(sandbox, monkeypatch):
    from core import config_loader, context_continuity, tool_dispatcher, llm_client
    monkeypatch.setattr(config_loader, "get_config", lambda: {"scheduler": {"owner_id": "reader"}})
    monkeypatch.setattr(tool_dispatcher, "get_tools_schema", lambda **kwargs: [{"function": {"name": "read_document"}}])
    monkeypatch.setattr(tool_dispatcher, "_is_tool_enabled", lambda _: True)
    doc = store("资料原文" * 3000)
    before = context_continuity.messages("reader", "char-test")
    context_continuity.acknowledge(before)
    await read_document_for_user("reader", "char-test", doc)
    monkeypatch.setattr(llm_client, "summarize_document_text", AsyncMock(return_value="持久概要内容"))
    await document_reading.summarize("reader", "char-test", doc)
    after = context_continuity.messages("reader", "char-test")
    assert not any(item["_layer"] == "10.6_pending_material" for item in after)
    assert "持久概要内容" in str(after) and "6000" in str(after)


@pytest.mark.asyncio
async def test_dispatcher_keeps_tail_of_document_page(sandbox, monkeypatch):
    from core import tool_dispatcher
    text = "begin " + "x" * 5800 + " END_OF_FIRST_PAGE " + "y" * 5000
    doc = store(text)
    # Exercise the real dispatcher and ToolResult/frame boundary, isolated from grants.
    monkeypatch.setattr(tool_dispatcher, "get_config", lambda: {"scheduler": {"owner_id": "reader"}})
    monkeypatch.setattr(tool_dispatcher, "_is_tool_enabled", lambda _: True)
    monkeypatch.setattr("core.self_management.policy.tool_allowed", lambda *args, **kwargs: True)
    outcome = await tool_dispatcher.execute_structured("read_document", {"document_id": doc}, "reader",
                                                      "reader", False, object(), char_id="char-test", origin="assistant_loop")
    assert outcome.status == "tool_executed"
    assert "END_OF_FIRST_PAGE" in frame_tool_message(outcome.result)
    assert "next_offset=6000" in outcome.result


@pytest.mark.asyncio
async def test_first_upload_background_finishes_multiple_batches_on_summary_route(sandbox, monkeypatch):
    from core import llm_client
    doc = store("片段正文" * 12000)
    small = AsyncMock(return_value="分段及全文概要")
    monkeypatch.setattr(llm_client, "summarize_document_text", small)
    document_reading.schedule_summaries("reader", "char-test")
    for _ in range(50):
        await asyncio.sleep(0)
        if ("reader", "char-test") not in document_reading._jobs:
            break
    state = library.get_record("reader", "char-test", doc)["overview"]
    assert state["status"] == "ready" and len(state["parts"]) == 5
    assert all(call.kwargs["char_id"] == "char-test" for call in small.call_args_list)


def test_http_upload_keeps_trusted_caption_separate_and_offers_full_document_reference(sandbox, monkeypatch):
    from admin.routers import chat as chat_router
    from core import pipeline_registry
    from types import SimpleNamespace
    from tests.test_chat_media import _client, _scoped
    captured = []

    async def chat(text, channel, **kwargs):
        captured.append((text, channel, kwargs))
        return {"reply": "ok", "turn_id": "fixture-turn", "msg_id": "fixture-turn"}

    monkeypatch.setattr(chat_router, "run_owner_chat_turn", chat)
    monkeypatch.setattr("core.config_loader.get_config", lambda: {"scheduler": {"owner_id": "reader"}})
    monkeypatch.setattr(pipeline_registry, "get", lambda: SimpleNamespace(_active_character_id="char-test"))
    client = _client(monkeypatch)
    _scoped(sandbox, "emt_chat", ["chat"])
    text = "文档原文" * 3000 + "尾页关键证据"
    response = client.post("/upload/ingest", headers={"Authorization": "Bearer emt_chat"},
                           files=[("files", ("story.md", text.encode(), "text/markdown"))],
                           data={"channel": "mobile", "message": "请读完全文"})
    assert response.status_code == 200
    projection, channel, kwargs = captured[0]
    assert "document_id=doc_" in projection and "next_offset" in projection
    assert kwargs["trusted_user_text"] == "请读完全文" and channel == "mobile"
    doc = library.search("reader", "char-test")[0]["document_id"]
    assert "尾页关键证据" in library.read("reader", "char-test", doc, offset=12000)["content"]


@pytest.mark.asyncio
async def test_document_summary_uses_existing_summary_route_and_frozen_character(monkeypatch):
    from core import llm_client
    from types import SimpleNamespace
    model = SimpleNamespace(prompt_style="narrative")
    routes, requests = [], []

    def resolve(category, *, char_id=None):
        routes.append((category, char_id))
        return model

    async def execute(**kwargs):
        requests.append(kwargs)
        prepared = kwargs["prepare"](model)
        assert prepared.messages[-1]["content"] == "末段证据"
        return SimpleNamespace(ok=True, value=SimpleNamespace(assistant_text="概要"))

    monkeypatch.setattr(llm_client, "get_model_client", resolve)
    monkeypatch.setattr(llm_client, "execute_create", execute)
    assert await llm_client.summarize_document_text("末段证据", char_id="frozen-char") == "概要"
    assert routes == [("summary", "frozen-char")]
    assert requests[0]["call_category"] == "summary" and requests[0]["char_id"] == "frozen-char"


@pytest.mark.asyncio
async def test_search_keeps_all_document_identities_and_image_digests(sandbox):
    from core.tools.character_recall import search_documents_for_user
    expected = []
    for index in range(8):
        digest = str(index) * 64
        expected.append(library.store_upload(uid="reader", char_id="char-test", filename="long-name-" * 15 + ".png",
                                             media_type="image/png", sha256=digest, searchable_text="描述" * 100,
                                             source="upload_image"))
    safe = to_tool_result(await search_documents_for_user("reader", "char-test")).safe_summary
    assert len(safe) < 5000 and all(identity in safe for identity in expected)
    assert "sha256=" + "7" * 64 in safe and "image/png" in safe
